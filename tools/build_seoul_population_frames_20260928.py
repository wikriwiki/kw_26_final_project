"""Read official Seoul survey workbooks and census CSV without editing them.

Targets contain population covariates only. Policy outcome benchmarks never
enter this builder. Household income is attached to resident-person records;
it is not individual salary or an existing persona spending decile.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
import hashlib
import json
import math
from pathlib import Path
import re
import zipfile

ROOT = Path(__file__).resolve().parents[1]
NS = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
ROW_TAG = "{" + NS['x'] + "}row"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def evidence(path):
    return {"path": str(path.resolve()), "sha256": sha(path)}


def write(path, payload):
    body = json.dumps(payload, ensure_ascii=False, indent=2) + '\n'
    if path.exists() and path.read_text(encoding='utf-8') != body:
        raise ValueError(f'frozen population evidence would change: {path}')
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(body, encoding='utf-8')


def workbook_labels(path, codebook=False):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0 if codebook else 1]
    labels = {str(r[0]): list(r) for r in sheet.iter_rows(values_only=True) if r[0] is not None}
    workbook.close()
    return labels


def selected_rows(path, required):
    """Stream only requested values from sheet 1; never evaluate XLSX formulas."""
    from lxml import etree
    with zipfile.ZipFile(path) as archive:
        shared = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            xml = etree.fromstring(archive.read('xl/sharedStrings.xml'))
            shared = [''.join(item.itertext()) for item in xml]
        columns = None
        with archive.open('xl/worksheets/sheet1.xml') as stream:
            for _, row in etree.iterparse(stream, events=('end',), tag=ROW_TAG):
                values = {}
                for cell in row:
                    coordinate = re.match(r'[A-Z]+', cell.get('r', ''))
                    if not coordinate:
                        continue
                    col = coordinate.group(0)
                    if columns is not None and col not in columns:
                        continue
                    raw = cell.find('x:v', NS)
                    value = raw.text if raw is not None else None
                    if cell.get('t') == 's' and value is not None:
                        value = shared[int(value)]
                    elif cell.get('t') == 'inlineStr':
                        node = cell.find('x:is', NS)
                        value = ''.join(node.itertext()) if node is not None else None
                    elif value is not None:
                        value = float(value)
                        if value.is_integer():
                            value = int(value)
                    values[col] = value
                if columns is None:
                    columns = {col: str(value) for col, value in values.items() if str(value) in required}
                    if set(columns.values()) != set(required):
                        raise ValueError(f'{path.name}: missing columns {set(required)-set(columns.values())}')
                else:
                    yield {label: values.get(col) for col, label in columns.items()}
                row.clear()
                while row.getprevious() is not None:
                    del row.getparent()[0]


def survey_income(source, out, year):
    suffix = '_data.xlsx' if year == 2021 else '_combined.xlsx'
    members = source / 'extracted' / f'{year}_members{suffix}'
    households = source / 'extracted' / f'{year}_households{suffix}'
    codebook = source / 'extracted' / f'{year}_households_codebook.xlsx' if year == 2021 else households
    member_book = source / 'extracted' / f'{year}_members_codebook.xlsx' if year == 2021 else members
    labels = workbook_labels(codebook, codebook=year == 2021)
    income_key = 'QAQ1' if year == 2020 else 'AAQ1' if year == 2021 else 'AQ1'
    assert '가구소득' in str(labels[income_key][2])
    choices = {}
    for line in str(labels[income_key][3]).splitlines():
        match = re.match(r'(\d+)\.(.+)', line)
        if match:
            choices[int(match.group(1))] = match.group(2).strip()
    questionnaire_evidence = []
    conflict_resolution = None
    if year == 2021:
        from lxml import etree
        # Published codebook has 20 labels, but the official questionnaire and
        # released data both use 21. Verify the original questionnaire instead
        # of recoding the 21st category as missing or borrowing another year.
        questionnaire = source / 'extracted/2021_household_questionnaire.hwpx'
        with zipfile.ZipFile(questionnaire) as package:
            sections = [etree.fromstring(package.read(name)) for name in package.namelist()
                        if name.startswith('Contents/section') and name.endswith('.xml')]
        text = ' '.join(''.join(node.itertext()) for section in sections for node in section.iter()
                        if node.tag.endswith('}t'))
        assert len(choices) == 20 and '900~950만원 미만' in text and '950~1,000만원 미만' in text
        assert '1,000만원 이상' in text and '2020.9.1.~2021.8.31.' in text
        choices[19], choices[20], choices[21] = '900~950만원 미만', '950~1,000만원 미만', '1,000만원 이상'
        questionnaire_evidence = [evidence(questionnaire)]
        conflict_resolution = '공개 코드북 20개 표기와 raw 21개 코드 충돌: 같은 2021 공식 조사표 AQ1의 21구간을 사용. 19=900~950, 20=950~1000, 21=1000만원 이상.'
    expected_bins = {2019:19, 2020:19, 2021:21, 2024:21, 2025:21}[year]
    assert len(choices) == expected_bins
    linked = {}
    id_key = '일련번호' if year >= 2024 else 'ID'
    member_id = 'id' if year == 2019 else id_key
    sex_key = 'sq1_2' if year == 2019 else 'SQ1_2'
    birth_key = 'sq1_3' if year == 2019 else 'SQ1_3'
    gu_key = 'gu' if year == 2019 else '자치구' if year == 2025 else 'GU'
    observed_member_income_key = 'aq1' if year == 2019 else 'AQ1' if year >= 2024 else None
    for row in selected_rows(households, [id_key, income_key]):
        if row[id_key] in linked:
            raise ValueError('duplicate household link ID')
        linked[row[id_key]] = row[income_key]
    fields = [member_id, sex_key, birth_key, 'wtb1', gu_key]
    if year not in [2019, 2020]:
        fields.append('DEW2')
    if observed_member_income_key:
        fields.append(observed_member_income_key)
    counts, joint = Counter(), Counter()
    missing, rejected_age, raw_rows, joined = 0, 0, 0, 0
    total_weight = 0.0
    for row in selected_rows(members, fields):
        raw_rows += 1
        income = linked.get(row[member_id])
        if observed_member_income_key and row[observed_member_income_key] != income:
            raise ValueError('member and household income labels differ')
        if income not in choices:
            missing += 1
            continue
        joined += 1
        if year in [2019, 2020]:
            # The source publishes birth year, not birthday; the convention is
            # explicit and cannot be claimed as exact age on an experiment day.
            age = year - int(row[birth_key])
            age_band = str((age // 10) * 10) + '대' if age < 60 else '60세 이상'
            adult = age >= 20
        else:
            code = int(row['DEW2'])
            age_band = {1:'10대',2:'20대',3:'30대',4:'40대',5:'50대',6:'60세 이상'}[code]
            adult = code >= 2
        if not adult:
            rejected_age += 1
            continue
        sex = {1:'M', 2:'F'}[int(row[sex_key])]
        weight = float(row['wtb1'])
        if not math.isfinite(weight) or not weight > 0:
            raise ValueError('nonfinite/nonpositive member expansion weight')
        gu = str(row[gu_key])
        band = choices[int(income)]
        counts[band] += weight
        joint[(gu, sex, age_band, band)] += weight
        total_weight += weight
    if missing:
        raise ValueError(f'{year}: {missing} unlinked/invalid income records')
    payload = {
        'schema':'seoul_survey_resident_income_distribution_v1', 'reference_year':year,
        'population_unit':'resident_person', 'age_scope':'20세 이상',
        'income_definition':'해당 시민이 속한 가구의 월평균 총소득 구간(개인 근로소득 아님)',
        'original_household_rows':len(linked), 'original_member_rows':raw_rows,
        'income_linked_member_rows':joined, 'adult_member_rows':joined-rejected_age,
        'adult_weight_total':total_weight, 'missing_income_rows':missing,
        'weight_field':'wtb1: 가구원 원가중치',
        'age_definition':'DEW2의 원문 연령구간' if year not in [2019,2020] else f'{year}−출생연도, 생일 없는 연말 나이 근사',
        'proportions':{name:counts[name]/total_weight for name in choices.values()},
        'weighted_counts':dict(counts),
        'conditional_income_weighted_cells':[
            {'gu':gu,'sex':sex,'age_band':age,'income_band':income,'weighted_count':amount}
            for (gu,sex,age,income),amount in sorted(joint.items())],
        'source_evidence':[evidence(p) for p in sorted({members,households,codebook,member_book})],
        'source_url':'https://data.seoul.go.kr/dataList/OA-15564/F/1/datasetView.do',
        'limitations':['공개 소득 원자료의 공간 식별자는 자치구이며 행정동별 소득 결합분포는 공개되지 않음.',
                       '구간 내 정확한 소득금액과 최상위 열린 구간의 상한을 임의로 만들지 않음.',
                       '행정동별 소득 할당에는 소지역 내 조건부 동일성 등의 가정과 독립적인 감사를 추가해야 함.'],
        'policy_outcome_used':False, 'uses_spending_decile_as_income':False}
    if questionnaire_evidence:
        payload['source_evidence'].extend(questionnaire_evidence)
        payload['codebook_conflict_resolution'] = conflict_resolution
        payload['income_reference_period'] = '2020-09-01~2021-08-31 세전 월평균 총 가구소득'
    path = out / f'income_distribution_{year}.json'
    write(path,payload)
    print(json.dumps({'year':year,'income_source':str(path),'adult_members':payload['adult_member_rows'],
                      'income_bins':len(counts),'sha256':sha(path)},ensure_ascii=False),flush=True)
    return path,payload


def frame(out, csv_path, year, income_path, income, quarter):
    with csv_path.open(encoding='cp949',newline='') as stream:
        records=[r for r in csv.DictReader(stream) if r['기준_년분기_코드'] == quarter]
    if not records:
        return None
    assert len({r['행정동_코드'] for r in records}) == len(records)
    counts = {field:Counter() for field in ['sex','age_band','admin_dong']}
    joint=[]
    for row in records:
        dong=row['행정동_코드'];name=row['행정동_코드_명']
        for age in ['20','30','40','50','60_이상']:
            total=int(row[f'연령대_{age}_상주인구_수'])
            male=int(row[f'남성연령대_{age}_상주인구_수'])
            female=int(row[f'여성연령대_{age}_상주인구_수'])
            assert male+female == total
            label='60세 이상' if age=='60_이상' else age+'대'
            counts['age_band'][label]+=total
            counts['admin_dong'][dong]+=total
            counts['sex']['M']+=male;counts['sex']['F']+=female
            for sex,number in [('M',male),('F',female)]:
                joint.append({'admin_dong':dong,'admin_dong_name':name,'sex':sex,'age_band':label,'count':number})
    totals={sum(v.values()) for v in counts.values()};assert len(totals)==1
    population=totals.pop();targets={}
    for field,counter in counts.items():
        targets[field]={'proportions':{k:v/population for k,v in sorted(counter.items())},
                        'source':evidence(csv_path),'population_unit':'resident_person','reference_year':year,
                        'source_population_unit':'resident_person','source_reference_year':year,
                        'field_definition':f'서울시 상주인구 제공 CSV 기준분기 {quarter}, 20세 이상 시민의 {field}',
                        'source_reference_period':quarter,'source_url':'https://data.seoul.go.kr/dataList/OA-22183/S/1/datasetView.do'}
    targets['income_band']={'proportions':income['proportions'],'source':evidence(income_path),
                            'population_unit':'resident_person','reference_year':year,
                            'source_population_unit':'resident_person','source_reference_year':year,
                            'field_definition':income['income_definition']+'; 가구원 원가중치로 20세 이상 시민에게 연결',
                            'source_url':income['source_url']}
    payload={'schema':'population_calibration_frame_v1','reference_year':year,'population_unit':'resident_person',
             'age_scope':'20세 이상','targets':targets,'resident_count':population,
             'administrative_dong_count':len(records),'demographic_joint_counts':joint,
             'status':'source_targets_prepared_not_cohort_validated',
             'limitations':['4개 주변분포 목표이며 성별×연령×행정동×소득의 실측 4차원 결합분포를 관측한 것은 아님.',
                            '서울시 원자료의 기준분기 라벨을 사용했으며 분기값 반복/집계 갱신의 관측시점은 별도 확인 필요.',
                            '실행 Agent의 거주 anchor 코드와 행정동 경계, 정확한 연령, 소득 할당을 감사하기 전 새 실험에 사용할 수 없음.']}
    path=out/f'population_frame_{year}.json';write(path,payload)
    print(json.dumps({'year':year,'frame':str(path),'dongs':len(records),'adult_population':population,'sha256':sha(path)},ensure_ascii=False),flush=True)
    return path


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--source-dir',type=Path,required=True)
    parser.add_argument('--out',type=Path,default=ROOT/'output/population_matching_20260928')
    parser.add_argument('--years',type=int,nargs='+',default=[2020,2021,2025])
    parser.add_argument('--income-only',action='store_true')
    args=parser.parse_args();args.out.mkdir(parents=True,exist_ok=True)
    for year in args.years:
        path,income=survey_income(args.source_dir,args.out,year)
        if year in [2021,2025] and not args.income_only:
            frame(args.out,args.source_dir/'official_resident_admin_dong.csv',year,path,income,
                  '20213' if year==2021 else '20252')
    if args.income_only:
        return
    write(args.out/'source_status.json',{'schema':'population_source_status_v1',
          'income_source_years':[2020,2021,2025],'full_target_frame_years':[2021,2025],
          'current_population_matched_simulation_completed':False,
          'remaining':['2020년의 정확한 행정동별 등록인구 원문 확보',
                       '표본과 행정동 경계/거주 anchor 정합',
                       '실측 소득 분포에 근거한 합성시민 소득 할당 및 감사',
                       '4축 추출/가중 분포 관문 후 별도 사전등록 새 실행']})


if __name__=='__main__':
    main()
