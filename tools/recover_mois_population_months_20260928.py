"""Recover public MOIS monthly age CSVs with original administrative codes.

No KOSIS names are used to choose a correspondence: the optional source
comparison uses complete 66-cell population vectors and uniqueness checks.
Existing population frames, graph, pilots and model processes are unchanged.
"""
from __future__ import annotations

import calendar
import csv
import hashlib
import io
import json
import re
import shutil
import sys
import urllib.parse
import urllib.request
from collections import Counter,defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime,timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/population_frame_sources_20260928'
CROOT=Path('C:/Users/Administrator/Documents/kw26_a100_recovery_20260928/population_frame_sources')
URL='https://jumin.mois.go.kr/downloadCsvAge.do?searchYearMonth=month&xlsStats=3'
PERIODS={'P013':'202004','P016':'202006','P014':'202008','DISTANCING_2020':'202010','P012':'202109','P010':'202506'}
AGE_MAP={-1:'총인구수',**{a+5:f'{a}~{a+4}세'for a in range(0,100,5)},101:'100세 이상'}
SEX_MAP={'계':'계','남':'남자','여':'여자'}


def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()


def fetch(month):
    path=OUT/f'official_mois_age_{month}_national.source.csv'
    form={'sltOrgType':'1','sltOrgLvl1':'A','sltOrgLvl2':'','gender':'gender','sum':'sum','sltUndefType':'',
          'searchYearStart':month[:4],'searchMonthStart':month[4:],'searchYearEnd':month[:4],'searchMonthEnd':month[4:],
          'sltOrderType':'1','sltOrderValue':'ASC','sltArgTypes':'5','sltArgTypeA':'0','sltArgTypeB':'100','category':'month'}
    if not path.exists():
        request=urllib.request.Request(URL,data=urllib.parse.urlencode(form).encode(),headers={'User-Agent':'Mozilla/5.0','Referer':'https://jumin.mois.go.kr/ageStatMonth.do'})
        with urllib.request.urlopen(request,timeout=45)as response:
            raw=response.read()
            note={'url':URL,'form':form,'content_type':response.headers.get('Content-Type'),
                  'disposition':response.headers.get('Content-Disposition'),'captured_at':datetime.now(timezone.utc).isoformat()}
        path.write_bytes(raw)
        path.with_suffix('.request.json').write_text(json.dumps(note,ensure_ascii=False,indent=2),encoding='utf-8')
    raw=path.read_bytes()
    if b'\x00'in raw:raise ValueError('NUL in MOIS source')
    rows=list(csv.DictReader(io.StringIO(raw.decode('cp949'))))
    assert len(rows)>3000
    prefix=f'{month[:4]}년{month[4:]}월_'
    assert all(k=='행정구역'or k.startswith(prefix)for k in rows[0])
    values={}
    labels={}
    long=[]
    for row in rows:
        match=re.fullmatch(r'(.*?)\s*\((\d{10})\)',row['행정구역'])
        assert match
        code=match[2]
        if not code.startswith('11'):continue
        area=match[1].strip().split(maxsplit=2)
        label={'gu':area[1]if len(area)>=2 else'계','dong':area[2]if len(area)>=3 else'계'}
        labels[code]=label
        vector={}
        for sex,sexname in SEX_MAP.items():
            for age,agefield in AGE_MAP.items():
                field=f'{prefix}{sex}_{agefield}'
                value=int(row[field].replace(',',''))
                assert value>=0
                vector[(sexname,age)]=value
                long.append({'reference_month':month,'mois_admin_code10':code,'sigungu_nm':label['gu'],
                             'epmndn_nm':label['dong'],'sex_nm':sexname,'agegrd_cd':age,
                             'agegrd_nm':agefield,'popl_cnt':value})
            assert int(row[f'{prefix}{sex}_연령구간인구수'].replace(',',''))==vector[(sexname,-1)]
            assert sum(v for(s,a),v in vector.items()if s==sexname and a!=-1)==vector[(sexname,-1)]
        assert all(vector[('남자',a)]+vector[('여자',a)]==vector[('계',a)]for a in AGE_MAP)
        assert code not in values
        values[code]=vector
    dongs={c:v for c,v in values.items()if labels[c]['dong']!='계'}
    gus={c:v for c,v in values.items()if labels[c]['gu']!='계'and labels[c]['dong']=='계'}
    assert len(gus)==25 and len(dongs)in{425,426}
    for code,vector in gus.items():
        assert all(sum(v[key]for c,v in dongs.items()if c[:5]==code[:5])==vector[key]for key in vector)
    assert all(sum(v[key]for v in gus.values())==values['1100000000'][key]for key in values['1100000000'])
    csvpath=OUT/f'official_mois_age_seoul_{month}.long.from_csv.csv'
    with csvpath.open('w',encoding='utf-8-sig',newline='')as handle:
        writer=csv.DictWriter(handle,fieldnames=list(long[0]));writer.writeheader();writer.writerows(long)
    return month,path,csvpath,values,labels


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    OUT.mkdir(parents=True,exist_ok=True);CROOT.mkdir(parents=True,exist_ok=True)
    Hpath=OUT/'official_code_sources/KIKcd_H.20250701(말소코드포함).xlsx'
    H={str(r[0]):r for r in list(load_workbook(Hpath,read_only=True,data_only=True).active.iter_rows(values_only=True))[1:]if r[0]}
    graphpath=OUT/'graph_full_population_code_projection_20260928.json'
    graph=json.loads(graphpath.read_text(encoding='utf-8'))
    graphcodes={r['code']for r in graph['dongs']}
    with ThreadPoolExecutor(max_workers=3)as executor:
        retrieved=dict((r[0],r)for r in executor.map(fetch,PERIODS.values()))
    results=[]
    for policy,month in PERIODS.items():
        _,path,csvpath,values,labels=retrieved[month]
        dongs={c:v for c,v in values.items()if labels[c]['dong']!='계'}
        day=month+f'{calendar.monthrange(int(month[:4]),int(month[4:]))[1]:02d}'
        lifecycle_fail=[]
        for code in dongs:
            h=H.get(code)
            if not h or str(h[4])>day or(h[5]and day>=str(h[5])):lifecycle_fail.append(code)
        assert not lifecycle_fail
        original=OUT/f'official_resident_sex_age_dong_{month}.source.json'
        raw=json.loads(original.read_text(encoding='utf-8'))['dataTypeList']
        native=defaultdict(dict)
        for r in raw:
            if r['epmndn_nm']!='계':native[r['admdong_cd']][(r['sex_nm'],r['agegrd_cd'])]=r['popl_cnt']
        assert len(native)==len(dongs)
        reverse=defaultdict(list)
        keys=sorted(next(iter(dongs.values())))
        for code,vector in dongs.items():reverse[tuple(vector[k]for k in keys)].append(code)
        assert all(len(codes)==1 for codes in reverse.values())
        correspondences=[]
        for native_code,vector in native.items():
            assert set(vector)==set(keys)
            matches=reverse.get(tuple(vector[k]for k in keys),[])
            assert len(matches)==1
            correspondences.append({'native_datuhub_kosis_code10':native_code,'direct_mois_code10':matches[0],
                                    'complete_66cell_vector_equal':True,'unique_among_all_seoul_dongs':True,
                                    'name_used_for_match':False,'is_same_code':native_code==matches[0]})
        assert len({r['direct_mois_code10']for r in correspondences})==len(dongs)
        missing=[{'mois_code10':c,'dong':labels[c]}for c in dongs if c[:-2]not in graphcodes]
        row={'policy':policy,'reference_month':month,'source_code_namespace':'MOIS_admin10_actual_at_reference_month',
             'source_encoding':'cp949','derived_csv_encoding':'UTF-8-BOM','seoul_dongs':len(dongs),
             'seoul_total_population':values['1100000000'][('계',-1)],
             'all66cell_integer_sex_age_and_gu_city_sum_checks_pass':True,
             'all_dong_codes_official_H_lifecycle_active_at_month':True,
             'source_evidence':[{'path':str(p.relative_to(ROOT)),'sha256':sha(p),'bytes':p.stat().st_size}for p in[path,csvpath,path.with_suffix('.request.json'),original,Hpath,graphpath]],
             'native_source_correspondence_method':'Exact population vector identification between two official sources, not an officially published code crosswalk',
             'complete_native_vectors_match_MOIS':len(correspondences),'different_code_correspondences':[r for r in correspondences if not r['is_same_code']],
             'all_correspondences':correspondences,
             'current_graph_dong_code_identity_matches':len(dongs)-len(missing),'historical_codes_missing_current_graph':missing,
             'historical_boundary_and_full_candidate_distribution_gate_pass':False,
             'existing_frames_or_pilots_modified':False,
             'scope_note':'Prefer the direct MOIS source for future inputs; leave recovered native-source originals intact. Same code alone does not imply unchanged jurisdiction; apply documented split/rename relations and validate all target home cells separately.'}
        results.append(row)
        for p in[path,csvpath,path.with_suffix('.request.json')]:
            dest=CROOT/p.name;shutil.copyfile(p,dest);assert sha(p)==sha(dest)
    manifest={'schema':'official_mois_population_months_v1','generator':{'path':str(Path(__file__).relative_to(ROOT)),'sha256':sha(Path(__file__))},
              'page_url':'https://jumin.mois.go.kr/ageStatMonth.do',
              'population_definition':'Resident registration total includes resident, unknown-residence and overseas-national registrations; foreign nationals excluded. Month-end, five-year age bands.',
              'model_calls':0,'graph_modified':False,'results':results}
    out=OUT/'official_mois_population_months_manifest.json'
    out.write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    dest=CROOT/out.name;shutil.copyfile(out,dest);assert sha(out)==sha(dest)
    print(json.dumps([{k:r[k]for k in['policy','reference_month','seoul_dongs','seoul_total_population','current_graph_dong_code_identity_matches','different_code_correspondences']}for r in results],ensure_ascii=False,indent=2))
    print('MANIFEST',sha(out),'C/G PASS')


if __name__=='__main__':main()
