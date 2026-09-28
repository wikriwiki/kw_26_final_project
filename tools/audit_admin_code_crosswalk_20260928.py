"""Audit observed code identity and vintage; never assign by dong name alone.

Official MOIS H/MIX sources, preserved Seoul DataHub native-code populations,
and current read-only graph projections remain separate namespaces. No source,
graph, completed simulation, or population profile is changed.
"""
from __future__ import annotations

import calendar
import csv
import hashlib
import json
import shutil
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output/population_frame_sources_20260928'
CODE = OUT / 'official_code_sources'
C_BACKUP = Path('C:/Users/Administrator/Documents/kw26_a100_recovery_20260928/population_frame_sources')
H_PATH = CODE / 'KIKcd_H.20250701(말소코드포함).xlsx'
MIX_PATH = CODE / 'KIKmix.20250701(말소코드포함).xlsx'
GRAPH_PATH = OUT / 'graph_full_population_code_projection_20260928.json'
PERIODS = {'P013':'202004', 'P016':'202006', 'P014':'202008',
           'DISTANCING_2020':'202010', 'P012':'202109', 'P010':'202506'}
NOTICE = 'https://www.mois.go.kr/frt/bbs/type001/commonSelectBoardArticle.do?bbsId=BBSMSTR_000000000052&nttId=118483'


def digest(path): return hashlib.sha256(path.read_bytes()).hexdigest()


def evidence(path):
    return {'path': str(path.relative_to(ROOT)), 'sha256': digest(path), 'bytes': path.stat().st_size}


def date(value): return str(value or '')


def alive(row, date_):
    return row['created'] <= date_ and (not row['abolished'] or date_ < row['abolished'])


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    raw_h = list(load_workbook(H_PATH, read_only=True, data_only=True).active.iter_rows(values_only=True))
    h_rows = [{'code': str(r[0]), 'gu': r[2], 'name': r[3],
               'created': date(r[4]), 'abolished': date(r[5])}
              for r in raw_h[1:] if r[0] and str(r[0]).startswith('11') and r[3]]
    assert len({r['code'] for r in h_rows}) == len(h_rows)
    H = {r['code']: r for r in h_rows}
    graph = json.loads(GRAPH_PATH.read_text(encoding='utf-8'))
    g_dongs = {r['code']: r for r in graph['dongs']}
    assert len(g_dongs) == len(graph['dongs']) == 427
    current_rows = []
    for code8, row in g_dongs.items():
        # This representation is accepted only after each code, exact label,
        # district and official lifecycle is verified, not by truncation alone.
        code10 = code8 + '00'
        h = H.get(code10)
        ok = bool(h and alive(h, '20250701') and row['gu'] == h['gu'] and row['name'] == h['name'])
        current_rows.append({'graph_code8': code8, 'official_mois_code10': code10,
                             'graph_name': row['name'], 'graph_gu': row['gu'],
                             'official_record': h, 'identity_verified': ok})
    assert all(r['identity_verified'] for r in current_rows)
    # Exact rows from the legacy input are compared to the new official source;
    # source-vintage inference still must not bypass a historical boundary gate.
    old_path = ROOT / 'data/neo4j_load/admin/KIKcd_H.xlsx'
    old_rows = [tuple(r) for r in load_workbook(old_path, read_only=True, data_only=True).active.iter_rows(values_only=True)
                if r and str(r[0]).startswith('11')]
    official_seoul = [tuple(r) for r in raw_h[1:] if r and str(r[0]).startswith('11')]
    old_seoul_equal = set(old_rows) == set(official_seoul)
    source_evidence = [evidence(H_PATH), evidence(MIX_PATH), evidence(GRAPH_PATH),
                       evidence(old_path), evidence(ROOT/'scripts/neo4j_load/01_admin.py')]
    period_rows = []
    for policy, month in PERIODS.items():
        path = OUT / f'official_resident_sex_age_dong_{month}.source.json'
        source_evidence.append(evidence(path))
        source = json.loads(path.read_text(encoding='utf-8'))['dataTypeList']
        dongs = {r['admdong_cd']:{'gu':r['sigungu_nm'],'name':r['epmndn_nm']}
                 for r in source if r['epmndn_nm'] != '계'}
        day = month + f'{calendar.monthrange(int(month[:4]),int(month[4:]))[1]:02d}'
        rows = []
        for native10, row in dongs.items():
            h = H.get(native10)
            in_graph = native10.endswith('00') and native10[:-2] in g_dongs
            if h is None: status = 'native_statistical_code_absent_official_H'
            elif not alive(h, day): status = 'native_statistical_code_not_active_at_source_month'
            elif h['gu'] != row['gu']: status = 'official_code_district_conflict'
            elif not in_graph: status = 'official_historical_code_absent_current_graph'
            else: status = 'official_code_identity_same_current_graph'
            rows.append({'native_population_code10':native10, 'population_name':row['name'],
                         'population_gu':row['gu'], 'official_record':h,
                         'population_label_differs_from_official': bool(h and row['name'] != h['name']),
                         'status':status,
                         'candidate_graph_code8':native10[:-2] if status == 'official_code_identity_same_current_graph' else None,
                         'automatic_name_matching_used':False})
        counts = Counter(r['status'] for r in rows)
        period_rows.append({'policy':policy, 'reference_month':month,
                            'population_native_namespace':'SeoulDataHub_KOSIS_DT_1B04005N_native_10digit',
                            'source_dong_count':len(rows), 'status_counts':dict(counts),
                            'strict_whole_city_crosswalk_gate_pass':len(counts)==1 and 'official_code_identity_same_current_graph' in counts,
                            'population_label_differences':sum(r['population_label_differs_from_official'] for r in rows),
                            'rows':rows})
    agents = graph['agents']
    agent_counts = Counter()
    raw_counts = Counter()
    anchor_counts = Counter()
    mismatches = []
    for r in agents:
        raw = r['raw_residence_code']
        raw_counts[str(raw)] += 1
        for anchor in r['anchor_codes']: anchor_counts[str(anchor)] += 1
        agent_counts['agents'] += 1
        agent_counts['exact_age_present'] += isinstance(r['age'], int) and not isinstance(r['age'], bool)
        agent_counts['age20plus_exact'] += isinstance(r['age'], int) and r['age'] >= 20
        agent_counts['sex_M_or_F'] += r['sex'] in {'M','F'}
        agent_counts['one_home_anchor'] += len(r['anchor_ids']) == 1
        agent_counts['no_home_anchor'] += not r['anchor_ids']
        agent_counts['multiple_home_anchors'] += len(r['anchor_ids']) > 1
        agent_counts['raw_code_current_graph_supported'] += raw in g_dongs
        agent_counts['one_anchor_linked_same_code'] += len(r['anchor_codes']) == 1 and r['anchor_codes'] == r['linked_codes']
        match = r['anchor_codes'] == [raw] and r['linked_codes'] == [raw]
        agent_counts['raw_anchor_linked_all_agree'] += match
        if not match:
            mismatches.append({k:r[k] for k in ['aid','raw_residence_code','anchor_ids','anchor_codes','linked_codes']})
    local_mapping = ROOT/'data/neo4j_load/admin/adm_code_mapping.csv'
    source_evidence.append(evidence(local_mapping))
    legacy = list(csv.DictReader(local_mapping.open(encoding='utf-8-sig')))
    legacy_conflicts = []
    for r in legacy:
        h = H.get(str(r['행안부코드'])+'00')
        if not h or r['행정동명'] != h['name']:
            legacy_conflicts.append({'legacy':r,'official_mois':h})
    mixed_rows = list(load_workbook(MIX_PATH, read_only=True, data_only=True).active.iter_rows(values_only=True))[1:]
    law_to_admin = defaultdict(set)
    admin_to_law = defaultdict(set)
    history_special = []
    for r in mixed_rows:
        if not r[0] or not str(r[0]).startswith('11') or not r[3] or not r[4]: continue
        if not r[7]:
            law_to_admin[str(r[4])].add(str(r[0]))
            admin_to_law[str(r[0])].add(str(r[4]))
        if r[3] in ['용신동','용두동','신설동','상일동','상일제1동','상일제2동','일원2동','개포3동','수유제3동','수유3동']:
            history_special.append({'administrative_code10':str(r[0]),'gu':r[2],'administrative_name':r[3],
                                    'legal_code10':str(r[4]),'legal_name':r[5],
                                    'created':date(r[6]),'abolished':date(r[7])})
    obj = {'schema':'admin_code_crosswalk_audit_v1', 'created_at':datetime.now(timezone.utc).isoformat(),
           'generator':evidence(Path(__file__)), 'source_evidence':source_evidence,
           'graph_modified':False,'model_calls':0,
           'official_source_notice':NOTICE,
           'overall_new_citywide_sampling_gate_pass':False,
           'current_graph_namespace':{'name':'MOIS_administrative10_without_terminal00',
                                      'source_version':'20250701', 'graph_dongs':427,
                                      'identity_verified':427,'all_exact_name_district_code_lifecycle_pass':True,
                                      'legacy_local_KIK_Seoul_rows_equal_official_source':old_seoul_equal,
                                      'not_SGIS_NSO_namespace':True,
                                      'rows':current_rows},
           'policy_source_crosswalks':period_rows,
           'full_runtime_agent_projection':{'path':str(GRAPH_PATH.relative_to(ROOT)),
                                             'sha256':digest(GRAPH_PATH),'counts':dict(agent_counts),
                                             'raw_code_counts':dict(raw_counts),'anchor_code_counts':dict(anchor_counts),
                                             'raw_anchor_mismatches':mismatches,
                                             'generated_income_tier_is_not_empirical_income':True},
           'legacy_adm_mapping_csv':{'rows':len(legacy),'official_label_or_code_conflicts':len(legacy_conflicts),
                                     'conflict_examples':legacy_conflicts[:15],
                                     'accepted_for_runtime_mapping':False,
                                     'reason':'No preserved official SGIS crosswalk provenance/vintage. The rows cannot bypass exact MOIS code/lifecycle audit.'},
           'legal_admin_relation':{'legal_codes':len(law_to_admin), 'admin_codes':len(admin_to_law),
                                    'legal_codes_covering_multiple_admins':sum(len(v)>1 for v in law_to_admin.values()),
                                    'admin_codes_covering_multiple_legal_dongs':sum(len(v)>1 for v in admin_to_law.values()),
                                    'one_to_one_conversion_not_valid_in_general':True,
                                    'special_history':history_special},
           'remaining_requirements':[
               'Resolve native statistical inactive-code aliases with original MOIS/KOSIS code correspondence; same name is insufficient.',
               'Historical boundary harmonization is required for missing former 용신동/상일동/일원2동. Never redistribute inhabitants by names or POI counts.',
               'Future frozen roster must use actual unique home anchors and exact age/sex; exclude or repair missing/mismatched identities before model calls.',
               'Observed income assignment must be separately frozen and audited; legacy generated income is not empirical income.',
               'Current graph/POI catalog is 2026 data. Code identity alone does not establish historical POI geography or a representative completed pilot.']}
    path = OUT/'admin_code_crosswalk_audit.json'
    path.write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')
    dest=C_BACKUP/path.name;shutil.copyfile(path,dest);assert digest(path)==digest(dest)
    print('AGENTS',json.dumps(dict(agent_counts),ensure_ascii=False))
    print('PERIODS',json.dumps([{k:r[k] for k in ['policy','reference_month','source_dong_count','status_counts','strict_whole_city_crosswalk_gate_pass']} for r in period_rows],ensure_ascii=False))
    print('LEGACY',len(legacy),len(legacy_conflicts),'LEGAL_MULTIPLE',sum(len(v)>1 for v in law_to_admin.values()))
    print('AUDIT',digest(path),'C/G PASS')


if __name__=='__main__':main()
