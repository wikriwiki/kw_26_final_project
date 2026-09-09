"""Normalize public policy/context data without altering agents or the simulation.

Uses openpyxl only to read the original public workbook. No spreadsheet is edited.
Inputs: fetch_covid_support_sources.py receipts and public source files.
Outputs: JSON reference inputs and explicit data-quality findings.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from collections import Counter
from datetime import date, datetime, timedelta
from html.parser import HTMLParser
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data/experiments/covid_support_2021"
START, END = date(2021, 8, 1), date(2021, 12, 31)


def write(name: str, value) -> None:
    (DATA / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Unknown source text encoding")


class Tables(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows = []
        self.row = None
        self.cell = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.row = []
        if tag in ("td", "th") and self.row is not None:
            self.cell = []

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self.cell is not None:
            self.row.append(" ".join("".join(self.cell).split()))
            self.cell = None
        if tag == "tr" and self.row is not None:
            self.rows.append(self.row)
            self.row = None


def prepare_rules() -> int:
    parser = Tables()
    parser.feed(decode((DATA / "sources/seoul_grant_rules.html").read_bytes()))
    thresholds = []
    for row in parser.rows:
        if len(row) != 4 or not re.fullmatch(r"\d+인(?: 맞벌이)?", row[0]):
            continue
        amounts = [None if x == "-" else int(x.replace(",", "")) for x in row[1:]]
        thresholds.append({"household_size": int(re.match(r"\d+", row[0])[0]), "multiple_earners": "맞벌이" in row[0], "employee_krw": amounts[0], "regional_krw": amounts[1], "mixed_krw": amounts[2]})
    assert len(thresholds) == 19, "Official eligibility table layout changed"
    assert thresholds[0]["employee_krw"] == 170000
    write("national_support_rules.json", {
        "source_id": "seoul_grant_rules", "runtime_integration": "not_implemented",
        "policy_name": "코로나 상생 국민지원금", "amount_per_eligible_person_krw": 250000,
        "announcement_on": "2021-08-30", "applications_online_from": "2021-09-06",
        "applications_offline_from": "2021-09-13", "applications_until": "2021-10-29",
        "use_until": "2021-12-31", "card_credit_lag_days_after_application": 1,
        "eligibility": {
            "insurance_reference_month": "2021-06", "exclude_long_term_care_premium": True,
            "comparison": "household_premium <= threshold",
            "household_size_cap_for_table": 10, "thresholds": thresholds,
            "excluded_if_property_tax_base_2020_krw_gt": 900000000,
            "excluded_if_financial_income_2020_krw_gt": 20000000,
            "regional_earner_min_comprehensive_income_2020_krw": 3000000,
            "household_membership_rules": "Use official household/dependent rules. Individual spending decile is not household insurance eligibility.",
            "exceptions_requiring_separate_handling": ["의료급여 수급자", "외국인·재외국민 자격", "가구구성·보험료 이의신청"]
        },
        "seoul_usage": {"region": "서울특별시 전역", "merchant_rule": "지역사랑상품권 가맹점", "year": 2021,
            "large_online_platform_payment_allowed": False,
            "delivery_order_paid_in_person_at_eligible_merchant_terminal_allowed": True,
            "merchant_registry_2021_downloaded": False},
        "agent_entitlement_manifest_present": False,
        "note": "Reference rules only. Not a legacy Policy JSON. Missing household data must not be replaced by fabricated observed values."
    })
    return len(thresholds)


def prepare_cases() -> dict:
    import openpyxl
    workbook = openpyxl.load_workbook(DATA / "sources/seoul_district_cases.xlsx", read_only=True, data_only=True)
    sheet = workbook.active
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    groups = [(i, header[i]) for i in range(3, sheet.max_column, 2)]
    assert len([g for _, g in groups if g.endswith("구")]) == 25
    records = []
    for raw in sheet.iter_rows(min_row=3, values_only=True):
        if not isinstance(raw[0], datetime) or not START <= raw[0].date() <= END:
            continue
        row = {"source_date": raw[0].date().isoformat(), "total_cases": int(raw[1]), "total_deaths": int(raw[2]), "areas": {g: {"cases": int(raw[i] or 0), "deaths": int(raw[i + 1] or 0)} for i, g in groups}}
        assert sum(x["cases"] for x in row["areas"].values()) == row["total_cases"], row["source_date"]
        assert sum(x["deaths"] for x in row["areas"].values()) == row["total_deaths"], row["source_date"]
        records.append(row)
    workbook.close()
    assert len(records) == (END - START).days + 1
    assert len({r["source_date"] for r in records}) == len(records)
    city_text = decode((DATA / "sources/seoul_city_cases.csv").read_bytes())
    city = {}
    for row in csv.reader(io.StringIO(city_text)):
        if row and re.fullmatch(r"\d{4}-\d\d-\d\d", row[0]):
            city[row[0]] = int(row[1].replace(",", ""))
    differences = [{"date": r["source_date"], "district_file_total": r["total_cases"], "city_file_total": city.get(r["source_date"])} for r in records if r["total_cases"] != city.get(r["source_date"])]
    aligned_differences = []
    for row in records:
        reference = (date.fromisoformat(row["source_date"]) + timedelta(days=1)).isoformat()
        row["city_reference_date"] = reference
        if row["total_cases"] != city.get(reference):
            aligned_differences.append({"source_date": row["source_date"], "city_reference_date": reference, "district_total": row["total_cases"], "city_total": city.get(reference)})
            row["quality_flags"] = ["city_total_differs_after_date_alignment"]
        else:
            row["quality_flags"] = []
    write("seoul_cases_daily.json", {"source_ids": ["seoul_district_cases", "seoul_city_cases_data"], "coverage_from": START.isoformat(), "coverage_until": END.isoformat(), "records": records,
        "date_semantics": "city_reference_date = district source_date + 1 day; remaining numeric differences are flagged. These are retrospective revised counts, not original as-published bulletins.",
        "runtime_availability_policy": "Recommended conservative assumption: city_reference_date < simulation_date at dawn. Publication timestamp is not supplied by the source. Record this assumption and perform lag sensitivity.",
        "other_areas": "기타 and 타시도 are retained for reconciliation, not assigned to the 25 Seoul districts."})
    return {"days": len(records), "districts": 25, "same_date_mismatch_count": len(differences), "city_reference_date_offset_days": 1, "aligned_mismatch_count": len(aligned_differences), "aligned_differences": aligned_differences}


def prepare_vaccination() -> dict:
    raw = (DATA / "sources/seoul_vaccination_data.txt").read_text(encoding="utf-8")
    # Public Sheet endpoint returns identifier keys and trailing commas. Parse as
    # data only; never execute the response as JavaScript or Python.
    normalized = re.sub(r'([{,]\s*)([A-Za-z_]\w*)\s*:', r'\1"\2":', raw)
    normalized = re.sub(r",\s*([}\]])", r"\1", normalized)
    payload = json.loads(normalized)
    assert len(payload["list"]) == payload["page"]["totalCount"], "Incomplete public data page"
    records = []
    for source in payload["list"]:
        value = source["S_VC_DT"]
        if not re.fullmatch(r"2021\.\d\d\.\d\d", value):
            continue
        day = date.fromisoformat(value.replace(".", "-"))
        if not START <= day <= END:
            continue
        def number(key):
            value = source.get(key)
            return None if value in (None, "", "null") else float(value.replace(",", ""))
        n = number("FIR_SUB")
        flags = []
        for count_key, rate_key in [("FIR_INC", "FIR_INC_RATE"), ("SCD_INC", "SCD_INC_RATE")]:
            count, rate = number(count_key), number(rate_key)
            if n and count is not None and rate is not None and abs(100 * count / n - rate) > 0.15:
                flags.append(f"{count_key}_vs_{rate_key}_inconsistent")
        records.append({"source_date": day.isoformat(), "raw": source, "quality_flags": flags})
    counts = Counter(r["source_date"] for r in records)
    duplicates = {d: n for d, n in counts.items() if n > 1}
    expected = {(START + timedelta(days=i)).isoformat() for i in range((END - START).days + 1)}
    missing = sorted(expected - set(counts))
    inconsistent = sum(bool(r["quality_flags"]) for r in records)
    write("seoul_vaccination_review.json", {"source_id": "seoul_vaccination_data", "runtime_usable": False, "reason": "Raw count/rate consistency and duplicate-date issues require reconciliation; preserve raw values, do not invent missing zeroes.", "duplicate_dates": duplicates, "missing_dates": missing, "records": sorted(records, key=lambda r: r["source_date"])})
    return {"records": len(records), "inconsistent_records": inconsistent, "duplicate_dates": duplicates, "missing_dates": missing, "runtime_usable": False}


def main() -> None:
    receipt = json.loads((DATA / "download_manifest.json").read_text(encoding="utf-8"))
    for row in receipt["sources"]:
        if row["status"] == "downloaded":
            assert hashlib.sha256((ROOT / row["path"]).read_bytes()).hexdigest() == row["sha256"], row["id"]
    result = {"rule_table_rows": prepare_rules(), "cases": prepare_cases(), "vaccination": prepare_vaccination(), "simulation_ready": False, "runtime_integration": "not_implemented"}
    write("data_quality.json", result)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
